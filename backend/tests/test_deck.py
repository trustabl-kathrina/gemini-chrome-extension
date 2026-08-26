"""Scene 6 — pitch deck: build_deck's .pptx + preview page and GET /pages/deck/{id}.pptx."""

import io
import re
import zipfile
from collections.abc import AsyncIterator, Iterator
from typing import Any

import httpx
import pytest
from fastapi import FastAPI

from dayflow.api.app import create_app
from dayflow.core.loader import MemoryConfigStore, default_config
from dayflow.core.pages import PageStore, set_default_pages
from dayflow.core.vault import MemoryBlobStore
from dayflow.tools.browser import BROWSER_TOOL_NAMES
from dayflow.tools.connectors import GITHUB_TOOLS
from dayflow.tools.deck import DECK_TOOLS, PPTX_MIME, build_deck, build_pptx, normalize_outline
from dayflow.tools.server import SERVER_TOOLS
from tests.test_api import FakeRunner

OUTLINE: list[dict[str, Any]] = [
    {"heading": "Problem", "bullets": ["Course files are scattered", "Deadlines are missed"]},
    {"heading": "Users", "bullets": ["KBTU students", "Anyone with a browser workflow"]},
    {"heading": "Solution", "bullets": ["A browser agent", "Skills as config"]},
    {"heading": "Demo", "bullets": ["Vault sync", "Lab solver"], "notes": "Show the panel transcript."},
    {"heading": "Architecture", "bullets": ["ADK on Cloud Run", "MV3 extension"]},
    {"heading": "Progress", "bullets": ["Scenes 1-3 green"]},
    {"heading": "Plan", "bullets": ["Ship scenes 4-6", "Chrome Web Store"]},
    {"heading": "Ask", "bullets": ["Feedback", "Beta testers"]},
]


@pytest.fixture
def pages() -> Iterator[PageStore]:
    store = PageStore(MemoryBlobStore(), public_url="https://brain.test")
    set_default_pages(store)
    yield store
    set_default_pages(None)


@pytest.fixture
def app(pages: PageStore) -> FastAPI:
    return create_app(store=MemoryConfigStore(), runner=FakeRunner(), pages=pages)  # type: ignore[arg-type]


@pytest.fixture
async def client(app: FastAPI) -> AsyncIterator[httpx.AsyncClient]:
    async with httpx.AsyncClient(transport=httpx.ASGITransport(app=app), base_url="http://brain.test") as c:
        yield c


def slides_in(data: bytes) -> list[str]:
    names = zipfile.ZipFile(io.BytesIO(data)).namelist()
    assert "ppt/presentation.xml" in names
    return [n for n in names if re.fullmatch(r"ppt/slides/slide\d+\.xml", n)]


# ---------- pptx ----------


def test_build_pptx_is_16_9_with_one_slide_per_section() -> None:
    from pptx import Presentation

    data = build_pptx("Dayflow", "diploma", normalize_outline(OUTLINE))
    assert len(slides_in(data)) == len(OUTLINE) + 1  # title slide + one per section
    prs: Any = Presentation(io.BytesIO(data))
    assert prs.slide_width is not None and prs.slide_height is not None
    assert round(prs.slide_width / prs.slide_height, 2) == 1.78
    texts = [sh.text_frame.text for slide in prs.slides for sh in slide.shapes if sh.has_text_frame]
    assert "Dayflow" in texts
    assert any("Course files are scattered" in t for t in texts)


def test_normalize_outline_accepts_the_shapes_gemini_writes() -> None:
    sections = normalize_outline(
        [
            {"title": "Problem", "points": "Files are scattered\n- Deadlines are missed"},
            {"heading": "Ask", "bullets": ["Feedback"], "speaker_notes": "30 seconds"},
            "Thanks",
        ]
    )
    assert [s.heading for s in sections] == ["Problem", "Ask", "Thanks"]
    assert sections[0].bullets == ["Files are scattered", "Deadlines are missed"]
    assert sections[1].notes == "30 seconds"
    assert sections[2].bullets == []


def test_normalize_outline_rejects_an_entry_without_a_heading() -> None:
    with pytest.raises(ValueError, match="no heading"):
        normalize_outline([{"bullets": ["a"]}])


# ---------- tool ----------


async def test_build_deck_returns_preview_and_pptx_urls(pages: PageStore) -> None:
    out = await build_deck("Dayflow — diploma project", OUTLINE, subtitle="dayflow-student/diploma")
    assert out["status"] == "success"
    assert out["slides"] == len(OUTLINE) + 1
    assert out["preview_url"] == f"https://brain.test/pages/deck/{out['id']}"
    assert out["pptx_url"] == f"{out['preview_url']}.pptx"
    stored = await pages.blobs.get(f"pages/deck/{out['id']}.pptx")
    assert stored is not None and len(slides_in(stored)) >= 6
    html = await pages.get("deck", out["id"])
    assert html is not None
    assert "Architecture" in html and out["pptx_url"] in html  # the preview links its own file
    assert html.count("<section") == len(OUTLINE) + 1


async def test_build_deck_refuses_a_deck_that_would_have_fewer_than_six_slides() -> None:
    out = await build_deck("Too short", OUTLINE[:3])
    assert out["status"] == "error"
    assert "at least 5" in out["error"]


async def test_build_deck_escapes_html_from_the_repo(pages: PageStore) -> None:
    outline = [*OUTLINE[:7], {"heading": "<script>alert(1)</script>", "bullets": ["a & b"]}]
    out = await build_deck("Deck", outline)
    html = await pages.get("deck", out["id"])
    assert html is not None
    assert "<script>alert(1)</script>" not in html
    assert "&lt;script&gt;" in html and "a &amp; b" in html


# ---------- route ----------


async def test_get_pptx_serves_the_bytes_without_auth(client: httpx.AsyncClient) -> None:
    out = await build_deck("Dayflow", OUTLINE)
    res = await client.get(f"/pages/deck/{out['id']}.pptx")
    assert res.status_code == 200
    assert res.headers["content-type"].startswith(PPTX_MIME)
    assert len(slides_in(res.content)) == len(OUTLINE) + 1


async def test_preview_page_is_served_next_to_the_pptx(client: httpx.AsyncClient) -> None:
    out = await build_deck("Dayflow", OUTLINE)
    res = await client.get(f"/pages/deck/{out['id']}")
    assert res.status_code == 200
    assert res.headers["content-type"].startswith("text/html")
    assert "Download .pptx" in res.text


async def test_unknown_or_malformed_deck_id_is_404(client: httpx.AsyncClient) -> None:
    assert (await client.get("/pages/deck/gLdFbn8HcQyEqUv3Zk5Yhw.pptx")).status_code == 404
    assert (await client.get("/pages/deck/short.pptx")).status_code == 404


# ---------- pack skill (what Gemini is told to do in scene 6) ----------


def pitch_deck_skill():
    skill = default_config().skill("pitch-deck")
    assert skill is not None
    return skill


def test_skill_declares_only_tools_that_exist() -> None:
    skill = pitch_deck_skill()
    known = BROWSER_TOOL_NAMES | set(GITHUB_TOOLS) | {f.__name__ for f in [*SERVER_TOOLS, *DECK_TOOLS]}
    assert set(skill.tools) <= known, set(skill.tools) - known
    assert {"get_file_contents", "build_deck", "download", "open_tab"} <= set(skill.tools)


def test_skill_steps_follow_the_scene_order() -> None:
    text = pitch_deck_skill().instructions
    steps = ["get_file_contents", "build_deck(", "download(url=<pptx_url>", "open_tab("]
    positions = [text.index(s) for s in steps]
    assert positions == sorted(positions), "steps must go: read the repo → build_deck → download → open_tab"
    # The harness looks for the .pptx under this vault path and for a deck of at least MIN_SECTIONS + 1 slides.
    assert "pitch/diploma-pitch.pptx" in text
    assert "8-section outline" in text and "at least 6 slides" in text
